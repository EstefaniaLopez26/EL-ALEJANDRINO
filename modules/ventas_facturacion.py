import streamlit as st
import pandas as pd
from datetime import datetime
import os
import tempfile
import io
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from fpdf import FPDF
from database.db_manager import registrar_auditoria

try:
    from openpyxl.drawing.image import Image as ExcelImage
except ImportError:
    ExcelImage = None

CODIGOS_PROMO = {
    "FERIA10": 10.0,
    "LECTOR15": 15.0,
    "BIENVENIDO15": 15.0,
    "ALEJANDRINO20": 20.0
}

class PDF(FPDF):
    def header(self):
        if os.path.exists("assets/logo.png"):
            self.image("assets/logo.png", 10, 8, 60)
            
        self.set_font("helvetica", "B", 16)
        self.set_text_color(20, 33, 61)
        self.cell(0, 8, "FACTURA COMERCIAL", 0, 1, "R")
        
        self.set_font("helvetica", "", 9)
        self.set_text_color(100, 100, 100)
        self.cell(0, 4, "Librería El Alejandrino", 0, 1, "R")
        self.cell(0, 4, "Sede: Duitama, Boyacá", 0, 1, "R")
        self.cell(0, 4, "elalejandrino1@gmail.com", 0, 1, "R")
        self.ln(12)
        
        self.set_draw_color(20, 33, 61)
        self.set_line_width(0.8)
        self.line(10, 35, 200, 35)
        self.ln(5)

    def footer(self):
        self.set_y(-30)
        self.set_font("helvetica", "I", 8)
        self.set_text_color(120, 120, 120)
        self.cell(0, 4, "Documento generado por el Sistema El Alejandrino", 0, 0, "C")

def generar_pdf_factura(factura_info, detalle_items):
    pdf = PDF()
    pdf.add_page()
    pdf.set_auto_page_break(auto=True, margin=15)

    pdf.set_font("helvetica", "B", 10)
    pdf.set_text_color(20, 33, 61)
    pdf.cell(95, 8, "DETALLES DE LA FACTURA", 0, 0, "L")
    pdf.cell(95, 8, "DATOS DEL CLIENTE", 0, 1, "L")
    
    pdf.set_font("helvetica", "", 9)
    pdf.set_text_color(0, 0, 0)
    pdf.cell(30, 6, "No. Factura:", 0, 0, "L")
    pdf.cell(65, 6, str(factura_info['numero']), 0, 0, "L")
    pdf.cell(30, 6, "Cliente:", 0, 0, "L")
    pdf.cell(65, 6, str(factura_info['cliente']), 0, 1, "L")
    
    pdf.cell(30, 6, "Fecha:", 0, 0, "L")
    pdf.cell(65, 6, str(factura_info['fecha']), 0, 0, "L")
    pdf.cell(30, 6, "Medio de Pago:", 0, 0, "L")
    pdf.cell(65, 6, str(factura_info['medio']), 0, 1, "L")
    
    pdf.ln(5)

    pdf.set_font("helvetica", "B", 9)
    pdf.set_fill_color(20, 33, 61)
    pdf.set_text_color(255, 255, 255)
    pdf.cell(25, 8, "ISBN", 1, 0, "C", True)
    pdf.cell(85, 8, "Descripción", 1, 0, "C", True)
    pdf.cell(18, 8, "Cant.", 1, 0, "C", True)
    pdf.cell(31, 8, "V. Unitario", 1, 0, "C", True)
    pdf.cell(31, 8, "Subtotal", 1, 1, "C", True)

    pdf.set_font("helvetica", "", 9)
    pdf.set_text_color(0, 0, 0)
    
    fill = False
    for item in detalle_items:
        pdf.set_fill_color(240, 240, 240) if fill else pdf.set_fill_color(255, 255, 255)
        pdf.cell(25, 8, str(item['isbn']), 1, 0, "C", fill=fill)
        pdf.cell(85, 8, str(item['titulo'][:42]), 1, 0, "L", fill=fill)
        pdf.cell(18, 8, str(item['cantidad']), 1, 0, "C", fill=fill)
        pdf.cell(31, 8, f"$ {item['precioUnitario']:,.0f}", 1, 0, "R", fill=fill)
        pdf.cell(31, 8, f"$ {item['subtotalLinea']:,.0f}", 1, 1, "R", fill=fill)
        fill = not fill

    pdf.ln(5)

    pdf.set_font("helvetica", "", 9)
    X_pos = 120
    ancho_lbl = 40
    ancho_val = 30
    
    pdf.set_fill_color(220, 220, 220)

    def fila_total(lbl, val, bold=False, fill_bg=False):
        pdf.set_x(X_pos)
        pdf.set_font("helvetica", "B" if bold else "", 10 if bold else 9)
        if bold:
            pdf.set_fill_color(20, 33, 61)
            pdf.set_text_color(255, 255, 255)
        pdf.cell(ancho_lbl, 8, lbl, 1, 0, "R", fill=bold or fill_bg)
        pdf.cell(ancho_val, 8, f"$ {val:,.2f}", 1, 1, "R", fill=bold or fill_bg)
        pdf.set_text_color(0, 0, 0)

    subtotal_real = factura_info['subtotal'] + factura_info['descuento']
    fila_total("Subtotal Bruto", subtotal_real, fill_bg=True)
    if factura_info['descuento'] > 0:
        fila_total("Descuento", -factura_info['descuento'], fill_bg=True)
    fila_total("IVA (19%)", factura_info['iva'], fill_bg=True)
    fila_total("TOTAL A PAGAR", factura_info['total'], bold=True)

    if not os.path.exists("facturas_pdf"):
        os.makedirs("facturas_pdf")
    nombre_archivo = os.path.join("facturas_pdf", f"{factura_info['numero']}.pdf")
    pdf.output(nombre_archivo)
    return nombre_archivo


def generar_excel_balance(df):
    wb = Workbook()
    ws = wb.active
    ws.title = "Balance_Ventas"

    header_fill = PatternFill(start_color="14213D", end_color="14213D", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True, size=11)
    title_font = Font(color="14213D", bold=True, size=16)
    subtitle_font = Font(color="555555", bold=True, size=12) 
    center_alignment = Alignment(horizontal="center", vertical="center")
    thin_border = Border(left=Side(style='thin', color='DDDDDD'), 
                         right=Side(style='thin', color='DDDDDD'), 
                         top=Side(style='thin', color='DDDDDD'), 
                         bottom=Side(style='thin', color='DDDDDD'))

    ws.row_dimensions[1].height = 45 
    
    if ExcelImage and os.path.exists("assets/logo.png"):
        try:
            img = ExcelImage("assets/logo.png")
            img.height = 50
            img.width = 150
            ws.add_image(img, 'A1')
        except Exception:
            pass

    ws.merge_cells('D1:J1')
    ws['D1'] = "LIBRERÍA EL ALEJANDRINO - BALANCE CONTABLE"
    ws['D1'].font = title_font
    ws['D1'].alignment = Alignment(horizontal="center", vertical="center")

    ws.merge_cells('D2:J2')
    ws['D2'] = f"Reporte Generado: {datetime.now().strftime('%Y-%m-%d')}"
    ws['D2'].font = subtitle_font
    ws['D2'].alignment = Alignment(horizontal="center", vertical="center")

    ws.append([]) 

    columnas = ['ID Factura', 'Cliente', 'Fecha Emisión', 'Subtotal', 'Descuento', 'IVA (19%)', 'Total Pagado', 'Medio Pago', 'Estado', 'Nota Crédito']
    ws.append(columnas)
    header_row = ws.max_row
    for cell in ws[header_row]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_alignment

    for index, row in df.iterrows():
        ws.append([
            row['numeroFactura'], row['cliente_nombre'], str(row['fechaEmision'])[:10], 
            row['subtotal'], row['descuento'], row['iva19'], 
            row['totalPagar'], row['medioPago'], row['estado'], row.get('notaCredito', '')
        ])
        current_row = ws.max_row
        for col_num, cell in enumerate(ws[current_row], 1):
            cell.border = thin_border
            cell.alignment = center_alignment
            if col_num in [4, 5, 6, 7]:
                cell.number_format = '"$"#,##0.00'

    column_widths = {'A': 15, 'B': 25, 'C': 15, 'D': 15, 'E': 15, 'F': 15, 'G': 15, 'H': 15, 'I': 12, 'J': 15}
    for col, width in column_widths.items():
        ws.column_dimensions[col].width = width

    ws.append([])
    totals_row = ws.max_row + 1
    ws.cell(row=totals_row, column=3, value="TOTALES:").font = Font(bold=True)
    ws.cell(row=totals_row, column=3).alignment = Alignment(horizontal="right", vertical="center")

    tot_sub = float(df['subtotal'].sum()) if not df.empty else 0.0
    tot_desc = float(df['descuento'].sum()) if not df.empty else 0.0
    tot_iva = float(df['iva19'].sum()) if not df.empty else 0.0
    tot_pag = float(df['totalPagar'].sum()) if not df.empty else 0.0

    for col_idx, val in zip([4, 5, 6, 7], [tot_sub, tot_desc, tot_iva, tot_pag]):
        cell = ws.cell(row=totals_row, column=col_idx, value=val)
        cell.font = Font(bold=True)
        cell.number_format = '"$"#,##0.00'
        cell.border = thin_border
        cell.fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
        cell.alignment = center_alignment

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


def vista_ventas_caja(conn, user):
    st.header("💰 Punto de Venta (POS) y Facturación")
    
    if 'orden_venta' not in st.session_state:
        st.session_state['orden_venta'] = []
    if 'descuento_pct' not in st.session_state:
        st.session_state['descuento_pct'] = 0.0

    df_cli = pd.read_sql_query("SELECT id, cedula, nombres, apellidos FROM clientes WHERE activo=1", conn)
    df_lib = pd.read_sql_query("SELECT isbn, titulo, precioVenta, stockDisponible FROM libros WHERE activo=1 AND stockDisponible > 0", conn)

    c1, c2 = st.columns([1.2, 1])
    with c1:
        st.subheader("1. Seleccionar Obras")
        with st.form("form_item_venta"):
            lib_opts = [f"{r['isbn']} - {r['titulo']} ($ {r['precioVenta']:,.0f} | Disp: {r['stockDisponible']})" for _, r in df_lib.iterrows()]
            item_sel = st.selectbox("Buscar Obra por ISBN o Título", lib_opts if lib_opts else ["Sin stock"])
            cant_v = st.number_input("Cantidad", min_value=1, value=1)
            if st.form_submit_button("➕ Agregar a la Orden"):
                if lib_opts:
                    isbn_v = item_sel.split(" - ")[0]
                    row_l = df_lib[df_lib['isbn'] == isbn_v].iloc[0]
                    if cant_v > row_l['stockDisponible']:
                        st.error(f"🚫 Solo hay {row_l['stockDisponible']} unidades disponibles.")
                    else:
                        sub_linea = round(float(row_l['precioVenta']) * cant_v, 2)
                        st.session_state['orden_venta'].append({
                            "isbn": isbn_v,
                            "titulo": row_l['titulo'],
                            "cantidad": int(cant_v),
                            "precioUnitario": float(row_l['precioVenta']),
                            "subtotalLinea": sub_linea
                        })
                        st.rerun()

        st.subheader("2. Convenios y Cupones Promocionales")
        col_p1, col_p2 = st.columns([2, 1])
        cod_input = col_p1.text_input("Código de Descuento (BIENVENIDO15, FERIA10...)")
        if col_p2.button("Validar Cupón"):
            cod_clean = cod_input.strip().upper()
            if cod_clean in CODIGOS_PROMO:
                st.session_state['descuento_pct'] = CODIGOS_PROMO[cod_clean]
                st.success(f"Descuento del {st.session_state['descuento_pct']}% aplicado.")
            else:
                st.error("Cupón inválido o expirado.")

    with c2:
        st.subheader("3. Liquidación y Factura")
        if not st.session_state['orden_venta']:
            st.info("Agrega libros en el panel izquierdo para iniciar la facturación.")
        else:
            df_orden = pd.DataFrame(st.session_state['orden_venta'])
            st.dataframe(df_orden[['isbn', 'titulo', 'cantidad', 'precioUnitario', 'subtotalLinea']], use_container_width=True, hide_index=True)
            
            subtotal_bruto = round(float(df_orden['subtotalLinea'].sum()), 2)
            valor_descuento = round(subtotal_bruto * (st.session_state['descuento_pct'] / 100.0), 2)
            subtotal_neto = round(subtotal_bruto - valor_descuento, 2)
            iva_19 = round(subtotal_neto * 0.19, 2)
            total_pagar = round(subtotal_neto + iva_19, 2)

            st.markdown(f"""
            - **Subtotal Bruto:** $ {subtotal_bruto:,.2f} COP
            - **Descuento ({st.session_state['descuento_pct']}%):** - $ {valor_descuento:,.2f} COP
            - **Base Gravable:** $ {subtotal_neto:,.2f} COP
            - **IVA (19%):** $ {iva_19:,.2f} COP
            - ### **Total a Pagar: $ {total_pagar:,.2f} COP**
            """)

            with st.form("form_cobrar_factura"):
                cli_opts = [f"{r['cedula']} - {r['nombres']} {r['apellidos']}" for _, r in df_cli.iterrows()]
                cli_venta = st.selectbox("Cliente", cli_opts)
                medio_pago = st.selectbox("Medio de Pago", ["Efectivo", "Tarjeta", "Transferencia / PSE"])
                monto_recibido = st.number_input("Valor Recibido ($ COP)", min_value=0.0, value=float(total_pagar), step=1000.0)
                
                if st.form_submit_button("✅ Emitir Factura Legal"):
                    if monto_recibido < total_pagar:
                        st.error("🚫 El monto recibido es inferior al total a cobrar.")
                    else:
                        ced_cli = cli_venta.split(" - ")[0]
                        nom_cli = cli_venta.split(" - ")[1]
                        id_cli = int(df_cli[df_cli['cedula'] == ced_cli].iloc[0]['id'])
                        fecha_emision = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
                        
                        c = conn.cursor()
                        c.execute("""
                            INSERT INTO ventas (cliente_id, empleado_id, fechaHora, subtotal, descuento, iva, total, estado)
                            VALUES (?, ?, ?, ?, ?, ?, ?, 'Pagada')
                        """, (id_cli, user['id'], fecha_emision, subtotal_neto, valor_descuento, iva_19, total_pagar))
                        venta_id = c.lastrowid
                        
                        for item in st.session_state['orden_venta']:
                            c.execute("""
                                INSERT INTO detalle_ventas (venta_id, isbn, titulo, cantidad, precioUnitario, subtotalLinea)
                                VALUES (?, ?, ?, ?, ?, ?)
                            """, (venta_id, item['isbn'], item['titulo'], item['cantidad'], item['precioUnitario'], item['subtotalLinea']))
                            c.execute("UPDATE libros SET stockDisponible = stockDisponible - ? WHERE isbn = ?", (item['cantidad'], item['isbn']))
                        
                        num_factura = f"FACT-{venta_id:05d}"
                        c.execute("""
                            INSERT INTO facturas (numeroFactura, venta_id, cliente_cedula, cliente_nombre, fechaEmision, subtotal, descuento, iva19, totalPagar, medioPago, montoRecibido, estado)
                            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 'Pagada')
                        """, (num_factura, venta_id, ced_cli, nom_cli, fecha_emision, subtotal_neto, valor_descuento, iva_19, total_pagar, medio_pago, monto_recibido))
                        conn.commit()
                        
                        cambio = round(monto_recibido - total_pagar, 2)
                        registrar_auditoria(conn, user['nombre'], user['rol'], "Facturación", "EMITIR_FACTURA", f"Emitió {num_factura} por ${total_pagar:,.2f}")
                        
                        factura_info = {
                            "numero": num_factura, "fecha": fecha_emision, "cliente": f"{nom_cli} ({ced_cli})",
                            "subtotal": subtotal_neto, "descuento": valor_descuento, "iva": iva_19,
                            "total": total_pagar, "medio": medio_pago
                        }
                        pdf_path = generar_pdf_factura(factura_info, st.session_state['orden_venta'])

                        st.session_state['ultima_factura'] = {
                            "info": factura_info, "cambio": cambio, "pdf_path": pdf_path
                        }
                        st.session_state['orden_venta'] = []
                        st.session_state['descuento_pct'] = 0.0
                        st.rerun()

            if st.button("🗑️ Limpiar Orden"):
                st.session_state['orden_venta'] = []
                st.session_state['descuento_pct'] = 0.0
                st.rerun()

        if 'ultima_factura' in st.session_state:
            f = st.session_state['ultima_factura']
            st.success(f"🎉 Factura **{f['info']['numero']}** generada. Cambio: **$ {f['cambio']:,.2f} COP**")
            
            with open(f['pdf_path'], "rb") as pdf_file:
                st.download_button(
                    label="📄 Descargar Factura PDF",
                    data=pdf_file,
                    file_name=f"{f['info']['numero']}.pdf",
                    mime="application/pdf",
                    use_container_width=True,
                    type="primary"
                )

def vista_anulaciones_reportes(conn, user):
    st.header("📊 Finanzas y Estadísticas Corporativas")
    
    df_v = pd.read_sql_query("SELECT * FROM facturas WHERE estado='Pagada'", conn)
    
    if not df_v.empty:
        df_v['fechaDT'] = pd.to_datetime(df_v['fechaEmision'])
        df_v['MesAnio'] = df_v['fechaDT'].dt.strftime('%Y-%m')
        
        # --- SELECTOR DE PERÍODO / MES (PARA FILTRAR HISTORIALES) ---
        col_f1, col_f2 = st.columns([1, 2])
        meses_disponibles = sorted(df_v['MesAnio'].unique(), reverse=True)
        mes_seleccionado = col_f1.selectbox("📅 Seleccionar Período (Mes/Año)", meses_disponibles)
        
        # Filtramos los datos para las tarjetas y gráficos
        df_filtrado = df_v[df_v['MesAnio'] == mes_seleccionado]
    else:
        df_filtrado = df_v
        mes_seleccionado = "Sin registros"

    # 1. Métricas de Impacto del Mes Seleccionado
    c1, c2, c3 = st.columns(3)
    c1.metric(f"Ingresos Netos ({mes_seleccionado})", f"$ {df_filtrado['subtotal'].sum():,.0f}" if not df_filtrado.empty else "$ 0")
    c2.metric("Impuestos (IVA)", f"$ {df_filtrado['iva19'].sum():,.0f}" if not df_filtrado.empty else "$ 0")
    c3.metric("Recaudo Total", f"$ {df_filtrado['totalPagar'].sum():,.0f}" if not df_filtrado.empty else "$ 0")
    st.markdown("---")

    t_rep, t_rank, t_anul = st.tabs(["📈 Balance y Gráficas", "🏆 Ranking Obras", "🚫 Notas Crédito"])
    
    with t_rep:
        st.subheader(f"Evolución de Ventas - Período {mes_seleccionado}")
        if not df_filtrado.empty:
            import altair as alt
            
            meses_es = {1: 'Ene', 2: 'Feb', 3: 'Mar', 4: 'Abr', 5: 'May', 6: 'Jun', 
                        7: 'Jul', 8: 'Ago', 9: 'Sep', 10: 'Oct', 11: 'Nov', 12: 'Dic'}
            
            df_filtrado['FechaEtiqueta'] = df_filtrado['fechaDT'].apply(lambda x: f"{x.day} {meses_es[x.month]}")
            
            ventas_dia = df_filtrado.groupby(['FechaEtiqueta', df_filtrado['fechaDT'].dt.date])['totalPagar'].sum().reset_index()
            ventas_dia = ventas_dia.sort_values('fechaDT')
            
            # Gráfica de Área Suave Estilo Financiero
            base = alt.Chart(ventas_dia).encode(
                x=alt.X('FechaEtiqueta:N', sort=None, title='Día del Mes', axis=alt.Axis(labelAngle=0)),
                y=alt.Y('totalPagar:Q', title='Ventas Totales ($ COP)')
            )
            
            area = base.mark_area(
                line={'color': '#14213D', 'strokeWidth': 3},
                color=alt.Gradient(
                    gradient='linear',
                    stops=[alt.GradientStop(color='#14213D', offset=0),
                           alt.GradientStop(color='rgba(20, 33, 61, 0.1)', offset=1)],
                    x1=1, x2=1, y1=1, y2=0
                )
            ).properties(height=350)
            
            puntos = base.mark_circle(size=60, color='#D4AF37')
            
            chart = (area + puntos).interactive()
            st.altair_chart(chart, use_container_width=True)
        else:
            st.info("No hay registros de ventas para el período seleccionado.")
            
        st.subheader("Historial de Facturación del Período")
        st.dataframe(df_filtrado[['fechaEmision', 'numeroFactura', 'cliente_nombre', 'totalPagar', 'medioPago']], use_container_width=True, hide_index=True)
        
        if not df_v.empty:
            excel_data = generar_excel_balance(df_filtrado) # Exporta exactamente el mes seleccionado
            st.download_button(
                label=f"📊 Exportar Balance de {mes_seleccionado} (Excel)",
                data=excel_data,
                file_name=f"Balance_Financiero_{mes_seleccionado}.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                type="primary"
            )

    with t_rank:
        st.subheader("Top Obras Más Vendidas")
        df_top = pd.read_sql_query("""
            SELECT titulo, SUM(cantidad) as unidades_vendidas, SUM(subtotalLinea) as ingresos
            FROM detalle_ventas
            GROUP BY titulo
            ORDER BY unidades_vendidas DESC
        """, conn)
        
        if not df_top.empty:
            df_chart = df_top.set_index('titulo')['unidades_vendidas']
            st.bar_chart(df_chart, color="#D4AF37")
            
        st.dataframe(df_top, use_container_width=True, hide_index=True)
        csv_ranking = df_top.to_csv(index=False).encode('utf-8')
        st.download_button("📊 Exportar Ranking (Excel CSV)", data=csv_ranking, file_name="Ranking_Obras.csv", mime="text/csv")

    with t_anul:
        df_fact = pd.read_sql_query("SELECT * FROM facturas WHERE estado='Pagada' ORDER BY id DESC", conn)
        if df_fact.empty:
            st.info("No hay facturas pagadas registradas.")
        else:
            f_sel = st.selectbox("Seleccionar Factura para Anular", df_fact['numeroFactura'].tolist())
            row_f = df_fact[df_fact['numeroFactura'] == f_sel].iloc[0]
            with st.form("form_anular"):
                st.write(f"**Cliente:** {row_f['cliente_nombre']} | **Valor:** $ {row_f['totalPagar']:,.2f}")
                motivo = st.text_area("Motivo de la Devolución")
                if st.form_submit_button("Emitir Nota Crédito"):
                    if not motivo.strip():
                        st.error("Especifica un motivo.")
                    else:
                        nc = f"NC-{row_f['numeroFactura'].split('-')[1]}"
                        c = conn.cursor()
                        c.execute("UPDATE facturas SET estado='Anulada', motivoAnulacion=?, notaCredito=? WHERE numeroFactura=?", (motivo.strip(), nc, f_sel))
                        c.execute("UPDATE ventas SET estado='Anulada' WHERE id=?", (int(row_f['venta_id']),))
                        
                        df_items = pd.read_sql_query("SELECT isbn, cantidad FROM detalle_ventas WHERE venta_id=?", conn, params=(int(row_f['venta_id']),))
                        for _, item in df_items.iterrows():
                            c.execute("UPDATE libros SET stockDisponible = stockDisponible + ? WHERE isbn = ?", (int(item['cantidad']), item['isbn']))
                        conn.commit()
                        registrar_auditoria(conn, user['nombre'], user['rol'], "Finanzas", "NOTA_CREDITO", f"NC: {nc} por anulación {f_sel}")
                        st.success(f"Nota Crédito {nc} emitida con éxito.")
                        st.rerun()

        st.subheader("Historial General de Documentos")
        st.dataframe(pd.read_sql_query("SELECT numeroFactura, cliente_nombre, fechaEmision, totalPagar, estado, notaCredito FROM facturas ORDER BY id DESC", conn), use_container_width=True, hide_index=True)